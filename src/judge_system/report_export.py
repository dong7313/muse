"""Portable Excel export without a separate Node.js script."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font


def export_workbook(report_dir: Path, filename: str = "judgement.xlsx") -> Path:
    rows = json.loads((report_dir / "records.json").read_text(encoding="utf-8"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Records"
    columns = list(dict.fromkeys(key for row in rows for key in row))
    if columns:
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            values = [row.get(key) for key in columns]
            values = [json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for v in values]
            sheet.append(values)
            # Model-produced strings must remain text, never spreadsheet formulas.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    catalog = report_dir / "rubric_catalog.json"
    if catalog.exists():
        rubrics = workbook.create_sheet("Rubrics")
        rubrics.append(["JSON"])
        payload = json.loads(catalog.read_text(encoding="utf-8"))
        for item in payload if isinstance(payload, list) else [payload]:
            rubrics.append([json.dumps(item, ensure_ascii=False)])
    target = report_dir / Path(filename).name
    workbook.save(target)
    return target
