import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from judge_system.pipeline import export_excel_report
from openpyxl import load_workbook


def load_script(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class ReportsTest(unittest.TestCase):
    def test_portable_share_and_excel(self):
        share = load_script("share", "scripts/bench_evaluate/package_share.py")
        with tempfile.TemporaryDirectory() as tmp:
            run = Path(tmp) / "run"
            reports = run / "reports"
            reports.mkdir(parents=True)
            code = run / "code.py"
            code.write_text("result = None\n")
            rows = [{"task_name": "box", "model_label": "test", "code_path": str(code),
                     "llm_judge_summary": "=1+1", "sandbox_error": str(run), "sandbox_ok": True, "geometry_valid": True, "watertight": True, "component_count_match": True}]
            (reports / "records.json").write_text(json.dumps(rows))
            (reports / "rubric_catalog.json").write_text("[]")
            workbook = load_workbook(export_excel_report(Path(tmp), run, "test.xlsx"))
            sheet = workbook["Records"]
            self.assertEqual(sheet.cell(2, 4).data_type, "s")
            output = share.package_run(run, Path(tmp) / "share")
            html = (output / "index.html").read_text()
            self.assertNotIn(str(run), html)
            self.assertIn("assets/00000.py", html)
            self.assertTrue((output / "assets/00000.py").exists())
            self.assertTrue(output.with_suffix(".tgz").exists())
            import tarfile
            with tarfile.open(output.with_suffix(".tgz")) as archive:
                self.assertTrue(all(not item.uname and not item.gname for item in archive.getmembers()))

    def test_scoring_with_mock_response(self):
        judge = load_script("judge_entry", "scripts/rerun_llm_judge_parallel.py")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            case = root / "box"
            case.mkdir()
            (case / "task.md").write_text("Make a box")
            (case / "plan.md").write_text("One solid")
            image = case / "image.png"
            image.write_bytes(b"fixture")
            record = {"task_name": "box", "png_path": str(image), "svg_path": "box.svg"}
            with patch.object(judge, "judge_with_vlm", return_value={"overall_score_normalized": 0.6}):
                row = judge._judge_one(record, cleaned_root=root, rubric_root=root,
                                       model="test", api_key_env="TEST_KEY", base_url="http://localhost", timeout_seconds=1)
            self.assertEqual(row["llm_judge_score"], 0.6)
            self.assertEqual(row["llm_judge_error"], "")

    def test_cli_help_in_installed_environment(self):
        env = os.environ.copy()
        env.pop("PYTHONPATH", None)
        for relative in ("scripts/rerun_llm_judge_parallel.py", "scripts/bench_evaluate/package_share.py"):
            result = subprocess.run([sys.executable, str(ROOT / relative), "--help"],
                                    env=env, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)

    def test_published_markdown_case_without_reference_code(self):
        from judge_system import reverse_pipeline as pipeline
        from judge_system.config import ModelSpec
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            case = root / "box"
            case.mkdir()
            (case / "design_description.md").write_text("Make a box")
            (case / "evaluation_rubric.md").write_text("Box must be solid")
            with patch.dict(os.environ, {"TEST_MODEL_KEY": "fixture"}), \
                 patch.object(pipeline, "_load_model_specs", return_value=[ModelSpec("test", "openrouter", "test", "TEST_MODEL_KEY")]), \
                 patch.object(pipeline, "_run_one_benchmark_case", return_value=object()) as worker:
                records, rubrics = pipeline._run_benchmark(
                    raw_root=root, task_root=root, rubric_root=root, run_root=root / "out",
                    cases=None, limit=None, api_key="", api_key_env="TEST_MODEL_KEY",
                    timeout_seconds=1, alignment_timeout_seconds=1, drawcad_root=root,
                    validator_root=root, python_executable=Path(sys.executable), model_list=None,
                    alignment_model="test", alignment_base_url="http://localhost", run_alignment=False,
                    alignment_system_prompt="", include_rubric_in_prompt=False, max_workers=1)
            self.assertEqual(len(records), 1)
            self.assertEqual(worker.call_args.kwargs["task_path"].name, "design_description.md")
            self.assertEqual(rubrics[0]["rubric_markdown"], "Box must be solid")
